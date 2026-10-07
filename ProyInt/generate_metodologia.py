import os
import docx
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_justified_indented_paragraph(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.5)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    return p

def main():
    doc = Document()

    # Define base styles
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(12)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("Reporte de Metodología y Plan de Proceso de Desarrollo Web\nProyecto: AsmaSync")
    run_title.bold = True
    run_title.font.size = Pt(14)
    run_title.font.name = 'Arial'

    p_space = doc.add_paragraph()

    # Section 1
    p_h1 = doc.add_paragraph()
    run_h1 = p_h1.add_run("1. Metodología de Desarrollo")
    run_h1.bold = True
    run_h1.font.size = Pt(12)
    run_h1.font.name = 'Arial'
    
    add_justified_indented_paragraph(doc, "Para el desarrollo del proyecto AsmaSync (Sistema inteligente de monitoreo y predicción de crisis asmáticas), se ha seleccionado la metodología ágil Scrum. La elección de Scrum se fundamenta en la naturaleza dinámica del proyecto y en la necesidad de integrar de forma iterativa y progresiva el frontend desarrollado en Angular, el backend construido en FastAPI y los algoritmos predictivos correspondientes.")
    add_justified_indented_paragraph(doc, "Scrum permite gestionar la complejidad del sistema dividiendo el desarrollo en ciclos cortos de trabajo denominados Sprints. Esto facilita la adaptación a nuevos requerimientos médicos o técnicos, garantizando que el sistema ofrezca valor desde las primeras fases. Asimismo, fomenta la colaboración continua entre el equipo de desarrollo y los expertos del dominio de la salud, asegurando que las funcionalidades de monitoreo y alerta temprana se ajusten a las necesidades reales de los pacientes.")

    p_space2 = doc.add_paragraph()

    # Section 2
    p_h2 = doc.add_paragraph()
    run_h2 = p_h2.add_run("2. Plan de Proceso de Desarrollo Web (Basado en Scrum)")
    run_h2.bold = True
    run_h2.font.size = Pt(12)
    run_h2.font.name = 'Arial'

    add_justified_indented_paragraph(doc, "El proceso de desarrollo web para AsmaSync se estructura en base a las ceremonias y artefactos de Scrum, garantizando una entrega continua de componentes funcionales y seguros. El plan se divide en las siguientes fases:")

    # Fase 1
    p_sub1 = doc.add_paragraph()
    run_sub1 = p_sub1.add_run("Fase 1: Inicio y Planificación (Product Backlog)")
    run_sub1.bold = True
    run_sub1.font.size = Pt(12)
    run_sub1.font.name = 'Arial'
    
    add_justified_indented_paragraph(doc, "En esta fase se definen los requerimientos generales del sistema AsmaSync, conformando el Product Backlog. Se priorizan las historias de usuario relacionadas con la seguridad, el registro de pacientes, la ingesta de datos de monitoreo en tiempo real y la visualización del dashboard. Se establece la arquitectura base, incluyendo la configuración inicial del entorno de Angular y la estructura del API REST en FastAPI.")

    # Fase 2
    p_sub2 = doc.add_paragraph()
    run_sub2 = p_sub2.add_run("Fase 2: Ejecución de Sprints de Desarrollo")
    run_sub2.bold = True
    run_sub2.font.size = Pt(12)
    run_sub2.font.name = 'Arial'

    add_justified_indented_paragraph(doc, "El desarrollo se divide en Sprints iterativos. Al inicio de cada Sprint, se realiza una reunión de Sprint Planning para seleccionar las historias de usuario a implementar. Durante el Sprint, el equipo desarrolla las funcionalidades correspondientes, llevando a cabo reuniones diarias (Daily Scrum) para sincronizar actividades y resolver bloqueos. Las actividades técnicas incluyen la codificación segura del backend, la creación de componentes responsivos en el frontend y la integración de modelos de Machine Learning.")

    # Fase 3
    p_sub3 = doc.add_paragraph()
    run_sub3 = p_sub3.add_run("Fase 3: Pruebas, Revisión y Retrospectiva")
    run_sub3.bold = True
    run_sub3.font.size = Pt(12)
    run_sub3.font.name = 'Arial'

    add_justified_indented_paragraph(doc, "Al finalizar cada Sprint, se lleva a cabo el Sprint Review para demostrar las nuevas funcionalidades operativas a los interesados (stakeholders), como los módulos de autenticación o los gráficos de monitoreo en el dashboard. Posteriormente, se realiza una Sprint Retrospective para identificar oportunidades de mejora en el proceso de desarrollo y en la comunicación del equipo, aplicando estos aprendizajes en el siguiente ciclo iterativo.")

    # Fase 4
    p_sub4 = doc.add_paragraph()
    run_sub4 = p_sub4.add_run("Fase 4: Despliegue y Mantenimiento Continuo")
    run_sub4.bold = True
    run_sub4.font.size = Pt(12)
    run_sub4.font.name = 'Arial'

    add_justified_indented_paragraph(doc, "Una vez que se alcanza un incremento del producto suficientemente maduro, se procede al despliegue en entornos de producción. La metodología ágil permite que, incluso después del lanzamiento inicial, el proyecto AsmaSync continúe recibiendo actualizaciones iterativas para mejorar la precisión de los modelos de predicción de crisis asmáticas y para mantener actualizadas las medidas de ciberseguridad del sistema.")

    doc.save(r"c:\asmasync-dashboard\ProyInt\Metodologia_Plan_Proceso.docx")
    print("Document successfully created!")

if __name__ == '__main__':
    main()
