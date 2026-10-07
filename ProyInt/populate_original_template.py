import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def populate_template():
    tpl_path = r"c:\asmasync-dashboard\ProyInt\Plantilla TRL4.docx"
    doc = Document(tpl_path)

    # Reemplazo de marcadores e información en el documento original
    for p in doc.paragraphs:
        txt = p.text.strip()
        if "Apaterno Amaterno Nombre(s)" in txt:
            p.text = "" # limpiar
        elif "Proyecto X" in txt:
            p.text = "AsmaSync - Sistema Inteligente de Monitoreo y Predicción de Crisis Asmáticas"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Montserrat"
                r.font.size = Pt(14)
                r.font.bold = True
        elif "Febrero del 2024" in txt:
            p.text = "Agosto de 2026"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Montserrat"
                r.font.size = Pt(14)

    # Insertar los integrantes en la carátula original
    for i, p in enumerate(doc.paragraphs):
        if "INTEGRANTES:" in p.text:
            # Reemplazar los siguientes 4 párrafos
            integrantes = [
                "Cerecedo Florencia Eliezer Isaí",
                "González Cuevas Juan Pablo",
                "Peña Ruiz Emmanuel",
                "Serrano Montaño Jocelyn"
            ]
            for idx, name in enumerate(integrantes):
                if i + 1 + idx < len(doc.paragraphs):
                    doc.paragraphs[i + 1 + idx].text = name
                    doc.paragraphs[i + 1 + idx].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in doc.paragraphs[i + 1 + idx].runs:
                        r.font.name = "Montserrat"
                        r.font.size = Pt(14)
            break

    out_path = r"c:\asmasync-dashboard\ProyInt\Plantilla_TRL4_Llenado.docx"
    doc.save(out_path)
    print("Populated template saved to:", out_path)

if __name__ == "__main__":
    populate_template()
