import os
import docx
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_justified_indented_paragraph(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.5)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    return p

def add_bullet_point(doc, title, text):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run_title = p.add_run(title)
    if title:
        run_title.bold = True
        run_title_colon = p.add_run(": ")
        run_title_colon.bold = True
        run_title_colon.font.name = 'Arial'
        run_title_colon.font.size = Pt(12)
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(12)
    run_text = p.add_run(text)
    run_text.font.name = 'Arial'
    run_text.font.size = Pt(12)
    return p

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12 if level > 1 else 14)
    run.font.name = 'Arial'
    return p

def main():
    doc = Document()

    # Define base styles
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(12)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("Documento del Proyecto Integrador\nSistema Inteligente AsmaSync")
    run_title.bold = True
    run_title.font.size = Pt(16)
    run_title.font.name = 'Arial'

    p_space = doc.add_paragraph()

    # 1. Plan del proceso de desarrollo
    add_heading(doc, "1. Plan del Proceso de Desarrollo Web (Metodología Ágil Scrum)")
    
    add_justified_indented_paragraph(doc, "El desarrollo del sistema AsmaSync se gestiona a través de la metodología ágil Scrum, lo que permite flexibilidad, entrega iterativa de valor y adaptación constante a los requerimientos técnicos y de salud.")
    
    add_bullet_point(doc, "Fase de Inicio y Planificación", "Definición del Product Backlog con las historias de usuario clave (Landing page, Dashboard de monitoreo médico, App móvil y predicción de IA). Se define la arquitectura base y los entornos de trabajo.")
    add_bullet_point(doc, "Ejecución de Sprints", "Sprints de 2 semanas donde equipos multidisciplinarios implementan el frontend, backend y modelos predictivos. Se realizan Daily Scrums para coordinar los avances.")
    add_bullet_point(doc, "Pruebas y Revisión", "Al término de cada Sprint, se ejecutan pruebas de QA y de seguridad. Se hace el Sprint Review para mostrar los avances a los interesados y una Retrospectiva para identificar mejoras para el siguiente ciclo.")
    add_bullet_point(doc, "Despliegue y Mantenimiento", "Lanzamientos iterativos al entorno de producción de las aplicaciones (Astro, Angular, Flutter) y del backend, garantizando una alta disponibilidad del sistema.")

    doc.add_paragraph()

    # 2. Justificación de la Arquitectura
    add_heading(doc, "2. Justificación de la Arquitectura")
    add_justified_indented_paragraph(doc, "La arquitectura seleccionada se basa en un enfoque de múltiples capas y microservicios orientados a eventos. Esta decisión se justifica por la necesidad crítica de procesar altos volúmenes de información provenientes de dispositivos IoT y ejecutar modelos predictivos de Machine Learning sin bloquear o ralentizar la experiencia del usuario final en las interfaces de presentación.")
    add_justified_indented_paragraph(doc, "La Capa de Presentación se desacopla del backend mediante un API Gateway y controles estrictos de Autorización (JWT) y Rate Limiting en la Capa de Aplicación. En la Capa de Negocio, la inclusión de Apache Kafka permite una distribución de eventos asíncrona, asegurando que el ingreso de datos, las notificaciones y el cálculo de riesgo operen de forma independiente, robusta y tolerante a fallos. Por último, la Capa de Persistencia segrega los datos almacenando la información transaccional de usuarios en PostgreSQL y los grandes flujos de datos estructurados de los sensores en MongoDB.")

    doc.add_paragraph()

    # 3. Diagrama de la arquitectura
    add_heading(doc, "3. Diagrama de la Arquitectura Seleccionada")
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_path = r"c:\asmasync-dashboard\ProyInt\security_architecture.png"
    if os.path.exists(img_path):
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.0))
    else:
        p_img.add_run("[Imagen del diagrama no encontrada en la ruta especificada]")

    doc.add_paragraph()

    # 4. Justificación de los frameworks
    add_heading(doc, "4. Justificación de los Frameworks de Desarrollo Utilizados")
    
    add_justified_indented_paragraph(doc, "La selección de tecnologías responde a la necesidad de crear un ecosistema de aplicaciones veloz, escalable y mantenible:")
    
    add_bullet_point(doc, "Astro (Landing Page Web)", "Permite un renderizado estático ultrarrápido y posee una arquitectura basada en 'islas'. Es la herramienta óptima para la página de presentación pública del proyecto, ya que maximiza el SEO y minimiza los tiempos de carga inicial al enviar cero JavaScript no esencial al navegador.")
    add_bullet_point(doc, "Angular (Dashboard Web Médico)", "Su elección se justifica por ser un framework robusto, fuertemente tipado (TypeScript) y enfocado en componentes. Su arquitectura madura resulta indispensable para construir un panel de control médico complejo que requiera manejar estados asíncronos y flujos de datos en tiempo real de forma segura.")
    add_bullet_point(doc, "Flutter (Aplicación Móvil)", "Se seleccionó gracias a su capacidad de generar aplicaciones nativas tanto para iOS como para Android a partir de un código base único (Dart). Garantiza un rendimiento gráfico excelente de 60fps y acelera considerablemente el despliegue de la solución hacia los pacientes.")
    add_bullet_point(doc, "FastAPI (Backend Web)", "Framework de desarrollo API en Python de última generación. Su uso se justifica ampliamente por su soporte asíncrono y altísimo rendimiento, siendo el puente ideal para interactuar de forma transparente con bibliotecas de Inteligencia Artificial y Machine Learning como Scikit-learn.")

    doc.add_paragraph()

    # 5. Lista de características
    add_heading(doc, "5. Herramientas y Frameworks Utilizados (Características Reales)")
    
    frameworks = [
        ("Astro", "Web / Generador de Sitios", "4.x", "Renderizado híbrido, Islands architecture, Optimizado 100% para SEO y Core Web Vitals."),
        ("Angular", "Web / SPA", "17.x", "Arquitectura basada en componentes, inyección de dependencias avanzada, Reactividad con RxJS, TypeScript estricto."),
        ("Flutter", "Móvil (iOS, Android)", "3.x", "Motor de renderizado UI de alto rendimiento (Skia/Impeller), Lenguaje Dart, Hot Reload para desarrollo ágil."),
        ("FastAPI", "Backend / API", "0.100+", "Rendimiento a la par de NodeJS/Go, validación automática de datos con Pydantic, soporte asíncrono (ASGI)."),
        ("Python", "Lenguaje / Data Science", "3.11+", "Uso intensivo en validación de datos y microservicios, soporte base para los algoritmos ML (Scikit-learn)."),
        ("Apache Kafka", "Broker de Mensajería", "3.x", "Sistema distribuido de eventos, alto 'throughput', baja latencia para encolado asíncrono de alertas de salud."),
        ("PostgreSQL", "Base de Datos Relacional", "16.x", "Alta fiabilidad e integridad referencial (ACID), almacenamiento robusto para usuarios y configuraciones médicas."),
        ("MongoDB", "Base de Datos NoSQL", "7.x", "Orientado a documentos JSON, flexibilidad de esquemas, excelente rendimiento para la ingesta de series temporales (IoT)."),
        ("Firebase Cloud Messaging", "Servicios Externos (Push)", "V1 API", "Entrega altamente confiable y multiplataforma de notificaciones de alerta directamente a la aplicación móvil.")
    ]

    for name, platform, version, desc in frameworks:
        p = doc.add_paragraph(style='List Bullet')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run_name = p.add_run(f"{name} ")
        run_name.bold = True
        run_name.font.name = 'Arial'
        run_name.font.size = Pt(12)
        
        run_details = p.add_run(f"(Versión: {version} | Plataforma: {platform}) - {desc}")
        run_details.font.name = 'Arial'
        run_details.font.size = Pt(12)

    doc.save(r"c:\asmasync-dashboard\ProyInt\Documento_Proyecto_Integrador.docx")
    print("Document successfully created!")

if __name__ == '__main__':
    main()
