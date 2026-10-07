import os
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_shading(cell, color):
    shading_elm = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), color))
    cell._tc.get_or_add_tcPr().append(shading_elm)

def add_placeholder(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.0)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F3F4F6") # Light gray
    
    # Set row height
    tr = table.rows[0]._tr
    trHeight = parse_xml(r'<w:trHeight xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="1440" w:hRule="exact"/>')
    tr.get_or_add_trPr().append(trHeight)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Add vertical spacing
    p.paragraph_format.space_before = Pt(30)
    run = p.add_run(f"📷 [ ESPACIO PARA CAPTURA DE PANTALLA ]\n{text}")
    run.font.name = 'Arial'
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(107, 114, 128)
    
    p_space = doc.add_paragraph() # Spacing after table
    p_space.paragraph_format.space_after = Pt(10)

def add_justified_indented_paragraph(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.5)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    return p

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12 if level > 1 else 14)
    run.font.name = 'Arial'
    return p

def main():
    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(12)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("Reporte de Implementación de Mecanismos de Autenticación Remota\nProyecto Integrador: AsmaSync")
    run_title.bold = True
    run_title.font.size = Pt(16)
    
    doc.add_paragraph()

    # Introduction
    add_justified_indented_paragraph(doc, "El presente reporte detalla los mecanismos de autenticación remota implementados en los Web Services del ecosistema AsmaSync, garantizando la seguridad en la comunicación entre el frontend (Angular / Flutter) y el backend (FastAPI), así como la integración con servicios externos de terceros.")

    # Service 1: API REST Backend (FastAPI)
    add_heading(doc, "Web Service 1: API REST Principal (FastAPI Backend)")
    
    add_heading(doc, "1) Justificación de los mecanismos implementados", level=2)
    add_justified_indented_paragraph(doc, "Para la API principal del sistema, se implementó el mecanismo de autenticación basada en Tokens Web JSON (JWT) utilizando el flujo OAuth2 (OAuth2PasswordBearer). La justificación radica en que JWT es un protocolo 'stateless' (sin estado), lo que significa que el servidor no necesita guardar sesiones en memoria, facilitando enormemente la escalabilidad horizontal del backend. Además, el token viaja en las cabeceras HTTP de forma segura, previniendo vulnerabilidades y permitiendo que tanto la aplicación web como la aplicación móvil utilicen el mismo mecanismo estandarizado para consumir los endpoints de manera unificada.")

    add_heading(doc, "2) Configuración general aplicada", level=2)
    add_justified_indented_paragraph(doc, "La configuración del JWT en el backend en Python incluye la definición de una clave secreta privada (SECRET_KEY) protegida mediante variables de entorno en el archivo .env. Se utiliza el algoritmo criptográfico HS256 para la firma del token. Adicionalmente, se configuró un tiempo de expiración (ACCESS_TOKEN_EXPIRE_MINUTES) riguroso para los tokens emitidos. Cuando un usuario envía sus credenciales al endpoint de '/login', el sistema valida el hash de la contraseña en la base de datos PostgreSQL y, de ser correcto, emite el token JWT que el cliente debe adjuntar en la cabecera 'Authorization: Bearer <token>' para sus peticiones futuras.")

    add_heading(doc, "3) Ejemplificación (con capturas de pantalla)", level=2)
    add_justified_indented_paragraph(doc, "A continuación, se ejemplifica de manera práctica cómo se realiza la petición para obtener el token de autenticación remota, y posteriormente cómo se utiliza para acceder a un recurso protegido en el Web Service:")
    
    add_placeholder(doc, "Captura sugerida 1: Interfaz de Swagger UI (https://asthma-predictor-api.onrender.com/docs#/) lanzando una petición POST al endpoint de '/login' y recibiendo el código 200 con el 'access_token' en la respuesta JSON.")
    add_placeholder(doc, "Captura sugerida 2: Interfaz de Swagger UI o Postman enviando una petición GET a un endpoint protegido (ej. '/patients') mostrando el candado cerrado o la cabecera 'Authorization: Bearer' inyectada exitosamente.")

    # Service 2: Firebase Cloud Messaging
    add_heading(doc, "Web Service 2: Firebase Cloud Messaging (Notificaciones Push)")
    
    add_heading(doc, "1) Justificación de los mecanismos implementados", level=2)
    add_justified_indented_paragraph(doc, "Para la emisión de alertas predictivas hacia la aplicación móvil, el backend debe comunicarse con el Web Service externo de Firebase Cloud Messaging (FCM). El mecanismo implementado para ello es la autenticación Server-to-Server mediante OAuth2 con una Cuenta de Servicio (Service Account). Se justifica la aplicación de este método porque es el protocolo de mayor seguridad avalado por Google, garantizando que únicamente nuestro servidor central autorizado (FastAPI) tenga los privilegios necesarios para despachar notificaciones masivas o directas hacia los pacientes.")

    add_heading(doc, "2) Configuración general aplicada", level=2)
    add_justified_indented_paragraph(doc, "La configuración general se llevó a cabo generando un archivo JSON de credenciales privadas desde la consola de Firebase y enlazándolo de forma local en el entorno del backend mediante el SDK de Firebase Admin. Durante el arranque del servidor, el SDK extrae y procesa el archivo de la Cuenta de Servicio, generando internamente los tokens de acceso OAuth2 necesarios y autenticando de manera transparente y continua todas las peticiones salientes hacia la API de mensajería externa de Firebase.")

    add_heading(doc, "3) Ejemplificación (con capturas de pantalla)", level=2)
    add_justified_indented_paragraph(doc, "El siguiente ejemplo ilustra la configuración de la credencial de autenticación Server-to-Server y cómo el backend valida sus credenciales remitiendo el payload hacia el Web Service de FCM:")

    add_placeholder(doc, "Captura sugerida 1: Pantalla de la consola de Firebase (Project Settings > Service accounts) donde se visualiza el botón de generación de la clave privada.")
    add_placeholder(doc, "Captura sugerida 2: Fragmento del código en el backend donde se inicializa Firebase Admin con el archivo JSON, o una captura de la terminal mostrando los logs de una notificación enviada exitosamente.")

    doc.save(r"c:\asmasync-dashboard\ProyInt\Reporte_Autenticacion_WebServices.docx")
    print("Document successfully created!")

if __name__ == '__main__':
    main()
