import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_shading(cell, color):
    shading_elm = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), color))
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_borders_code(cell, border_color):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        r'<w:tcBorders {}><w:top w:val="none"/><w:left w:val="single" w:sz="36" w:space="0" w:color="{}"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>'
        .format(nsdecls('w'), border_color)
    )
    tcPr.append(borders)

def add_hyperlink(paragraph, url, text, color="2563EB", underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = parse_xml(r'<w:hyperlink r:id="{}" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>'.format(r_id))
    new_run = parse_xml(r'<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
    rPr = OxmlElement('w:rPr')
    
    # Set font to Segoe UI
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Segoe UI')
    rFonts.set(qn('w:hAnsi'), 'Segoe UI')
    rPr.append(rFonts)

    c = OxmlElement('w:color')
    c.set(qn('w:val'), color)
    rPr.append(c)

    if underline:
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rPr.append(u)

    new_run.append(rPr)
    text_node = OxmlElement('w:t')
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink

def apply_text_formatting(paragraph, text, base_font_name="Segoe UI", base_font_size=11, base_color=(31, 41, 55)):
    # Parse bold (**bold**), code (`code`), and links ([text](url))
    # A simple parser using regex
    parts = re.split(r'(\*\*.*?\*\*|`.*?`|\[.*?\]\(.*?\))', text)
    
    for part in parts:
        if not part:
            continue
        
        if part.startswith('**') and part.endswith('**'):
            # Bold text
            inner_text = part[2:-2]
            run = paragraph.add_run(inner_text)
            run.font.name = base_font_name
            run.font.size = Pt(base_font_size)
            run.font.color.rgb = RGBColor(*base_color)
            run.bold = True
        elif part.startswith('`') and part.endswith('`'):
            # Inline Code
            inner_text = part[1:-1]
            run = paragraph.add_run(inner_text)
            run.font.name = "Consolas"
            run.font.size = Pt(base_font_size - 1)
            run.font.color.rgb = RGBColor(17, 24, 39)
            # Give it a light gray shading or different style
        elif part.startswith('[') and '](' in part and part.endswith(')'):
            # Hyperlink
            match = re.match(r'\[(.*?)\]\((.*?)\)', part)
            if match:
                link_text, link_url = match.groups()
                add_hyperlink(paragraph, link_url, link_text)
        else:
            # Normal text
            run = paragraph.add_run(part)
            run.font.name = base_font_name
            run.font.size = Pt(base_font_size)
            run.font.color.rgb = RGBColor(*base_color)

def main():
    doc = Document()
    
    # Configuración de márgenes a 1 pulgada
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # --- PÁGINA DE PORTADA (PREMIUM) ---
    # Título Principal
    p_title_space = doc.add_paragraph()
    p_title_space.paragraph_format.space_before = Pt(120)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_title = p_title.add_run("AsmaSync")
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(38)
    run_title.font.color.rgb = RGBColor(30, 58, 138) # Navy Blue
    run_title.bold = True
    p_title.paragraph_format.space_after = Pt(2)
    
    p_subtitle = doc.add_paragraph()
    run_sub = p_subtitle.add_run("Plan de Implementación de Mecanismos de Seguridad")
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(18)
    run_sub.font.color.rgb = RGBColor(75, 85, 99) # Slate Grey
    run_sub.bold = False
    p_subtitle.paragraph_format.space_after = Pt(40)
    
    # Separador
    p_sep = doc.add_paragraph()
    p_sep_run = p_sep.add_run("―" * 40)
    p_sep_run.font.color.rgb = RGBColor(37, 99, 235) # Accent Blue
    p_sep.paragraph_format.space_after = Pt(80)
    
    # Información del documento
    p_info = doc.add_paragraph()
    run_info = p_info.add_run(
        "Actividad 2: Especificación de Principios de Codificación Segura\n"
        "Curso: Proyecto Integrador\n"
        "Fecha: Mayo de 2026\n"
        "Entorno: Angular 17 & FastAPI Backend\n"
        "Estatus del Documento: Aprobado"
    )
    run_info.font.name = "Segoe UI"
    run_info.font.size = Pt(11)
    run_info.font.color.rgb = RGBColor(107, 114, 128)
    p_info.paragraph_format.line_spacing = 1.3
    
    doc.add_page_break()
    
    # --- LEER ARCHIVO MARKDOWN ---
    md_path = "Plan_Implementacion_Seguridad.md"
    if not os.path.exists(md_path):
        # Intentar ruta absoluta
        md_path = r"c:\asmasync-dashboard\ProyInt\Plan_Implementacion_Seguridad.md"
        
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    in_code_block = False
    code_lines = []
    code_lang = ""
    
    # Omitimos el primer título si ya lo pusimos en la portada
    skip_first_title = True
    
    for line in lines:
        stripped = line.strip()
        
        # Manejo de bloques de código
        if stripped.startswith("```"):
            if in_code_block:
                # Terminar bloque de código
                in_code_block = False
                
                # Crear tabla para el bloque de código
                table = doc.add_table(rows=1, cols=1)
                table.autofit = False
                table.columns[0].width = Inches(6.5)
                cell = table.cell(0, 0)
                
                # Sombreado y bordes
                set_cell_shading(cell, "F3F4F6") # Light gray background
                set_cell_borders_code(cell, "2563EB") # Left border blue
                
                # Añadir líneas de código en la celda
                p_cell = cell.paragraphs[0]
                p_cell.paragraph_format.space_before = Pt(4)
                p_cell.paragraph_format.space_after = Pt(4)
                p_cell.paragraph_format.line_spacing = 1.15
                
                code_text = "".join(code_lines)
                run_code = p_cell.add_run(code_text)
                run_code.font.name = "Consolas"
                run_code.font.size = Pt(8.5)
                run_code.font.color.rgb = RGBColor(17, 24, 39)
                
                # Añadir un espacio después del bloque
                p_after = doc.add_paragraph()
                p_after.paragraph_format.space_before = Pt(0)
                p_after.paragraph_format.space_after = Pt(6)
                
                code_lines = []
            else:
                in_code_block = True
                code_lang = stripped[3:].strip()
            continue
            
        if in_code_block:
            code_lines.append(line)
            continue
            
        # Encabezados
        if stripped.startswith("#"):
            match = re.match(r"^(#+)\s*(.*)$", stripped)
            if match:
                level = len(match.group(1))
                title_text = match.group(2)
                
                if level == 1:
                    if skip_first_title:
                        skip_first_title = False
                        continue
                    p = doc.add_heading(level=1)
                    p.paragraph_format.space_before = Pt(24)
                    p.paragraph_format.space_after = Pt(8)
                    p.paragraph_format.keep_with_next = True
                    run = p.add_run(title_text)
                    run.font.name = "Segoe UI"
                    run.font.size = Pt(20)
                    run.font.color.rgb = RGBColor(30, 58, 138)
                    run.bold = True
                elif level == 2:
                    p = doc.add_heading(level=2)
                    p.paragraph_format.space_before = Pt(18)
                    p.paragraph_format.space_after = Pt(6)
                    p.paragraph_format.keep_with_next = True
                    run = p.add_run(title_text)
                    run.font.name = "Segoe UI"
                    run.font.size = Pt(15)
                    run.font.color.rgb = RGBColor(37, 99, 235)
                    run.bold = True
                elif level == 3:
                    p = doc.add_heading(level=3)
                    p.paragraph_format.space_before = Pt(12)
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.keep_with_next = True
                    run = p.add_run(title_text)
                    run.font.name = "Segoe UI"
                    run.font.size = Pt(12.5)
                    run.font.color.rgb = RGBColor(75, 85, 99)
                    run.bold = True
                else:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(10)
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.keep_with_next = True
                    run = p.add_run(title_text)
                    run.font.name = "Segoe UI"
                    run.font.size = Pt(11)
                    run.font.color.rgb = RGBColor(31, 41, 55)
                    run.bold = True
            continue
            
        # Imágenes
        if stripped.startswith("![") and "](" in stripped:
            # Es una imagen
            match = re.match(r"^!\[(.*?)\]\((.*?)\)$", stripped)
            if match:
                alt_text, img_name = match.groups()
                # Buscar imagen en la carpeta
                img_path = img_name
                if not os.path.exists(img_path):
                    img_path = os.path.join("ProyInt", img_name)
                if not os.path.exists(img_path):
                    img_path = os.path.join(r"c:\asmasync-dashboard\ProyInt", img_name)
                    
                if os.path.exists(img_path):
                    # Insertar imagen centrada
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_before = Pt(12)
                    p_img.paragraph_format.space_after = Pt(4)
                    
                    p_img.add_run().add_picture(img_path, width=Inches(5.8))
                    
                    # Epígrafe / Leyenda de imagen
                    p_caption = doc.add_paragraph()
                    p_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_caption.paragraph_format.space_before = Pt(0)
                    p_caption.paragraph_format.space_after = Pt(12)
                    run_caption = p_caption.add_run(f"Figura: {alt_text}")
                    run_caption.font.name = "Segoe UI"
                    run_caption.font.size = Pt(9)
                    run_caption.font.italic = True
                    run_caption.font.color.rgb = RGBColor(107, 114, 128)
                else:
                    # Fallback si no está la imagen física
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(6)
                    run_err = p.add_run(f"[Imagen no encontrada: {img_name} - {alt_text}]")
                    run_err.font.name = "Segoe UI"
                    run_err.font.italic = True
                    run_err.font.color.rgb = RGBColor(220, 38, 38)
            continue
            
        # Línea horizontal
        if stripped == "---":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run("❖   ❖   ❖")
            run.font.name = "Segoe UI"
            run.font.color.rgb = RGBColor(209, 213, 219)
            continue
            
        # Listas de viñetas
        if stripped.startswith("* ") or stripped.startswith("- "):
            list_text = stripped[2:]
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            apply_text_formatting(p, list_text)
            continue
            
        # Listas numeradas
        match_num = re.match(r"^(\d+)\.\s*(.*)$", stripped)
        if match_num:
            num, list_text = match_num.groups()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            # Añadir número manualmente con indentación
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            
            run_num = p.add_run(f"{num}.  ")
            run_num.font.name = "Segoe UI"
            run_num.font.size = Pt(11)
            run_num.bold = True
            run_num.font.color.rgb = RGBColor(31, 41, 55)
            
            apply_text_formatting(p, list_text)
            continue
            
        # Párrafos normales
        if stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.15
            apply_text_formatting(p, stripped)
            
    # Guardar documento
    output_path = "Plan_Implementacion_Seguridad.docx"
    doc.save(output_path)
    print(f"Documento guardado con éxito en: {output_path}")

if __name__ == "__main__":
    main()
