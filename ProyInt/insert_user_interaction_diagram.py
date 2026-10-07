import os
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def format_run(run, text, font_name="Montserrat", font_size=14, color_rgb=(51, 65, 85), bold=False, italic=False):
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(*color_rgb)
    run.bold = bold
    run.italic = italic

def format_p(p, text, font_size=14, color_rgb=(51, 65, 85), bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6):
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    p.text = ""
    run = p.add_run(text)
    format_run(run, text, font_size=font_size, color_rgb=color_rgb, bold=bold, italic=italic)
    return p

def insert_image_after(p_ref, img_path, caption_text, width_inches=5.8):
    if not os.path.exists(img_path):
        print(f"Warning: Image {img_path} not found.")
        return p_ref

    p_img_elem = parse_xml(f'<w:p {nsdecls("w")}/>')
    p_ref._p.addnext(p_img_elem)
    doc_p_img = docx.text.paragraph.Paragraph(p_img_elem, p_ref._parent)
    doc_p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc_p_img.paragraph_format.space_before = Pt(14)
    doc_p_img.paragraph_format.space_after = Pt(4)
    
    run_img = doc_p_img.add_run()
    run_img.add_picture(img_path, width=Inches(width_inches))

    p_cap_elem = parse_xml(f'<w:p {nsdecls("w")}/>')
    p_img_elem.addnext(p_cap_elem)
    doc_p_cap = docx.text.paragraph.Paragraph(p_cap_elem, p_ref._parent)
    doc_p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc_p_cap.paragraph_format.space_after = Pt(16)
    
    run_cap = doc_p_cap.add_run(caption_text)
    format_run(run_cap, caption_text, font_size=12, color_rgb=(100, 116, 139), italic=True)
    return doc_p_cap

def insert_user_diagram_to_docx():
    src_img = r"C:\Users\penar\.gemini\antigravity\brain\ac6baca2-e01f-4959-b6b9-d862274c5e0a\user_interaction_diagram_1785956488652.jpg"
    target_img = r"c:\asmasync-dashboard\ProyInt\user_interaction_diagram.jpg"

    if os.path.exists(src_img):
        shutil.copy(src_img, target_img)
        print(f"Copied generated diagram to {target_img}")

    doc_files = [
        r"c:\asmasync-dashboard\ProyInt\Plantilla_TRL4_Capas.docx",
        r"c:\asmasync-dashboard\ProyInt\Plantilla_TRL4_Extenso.docx",
        r"c:\asmasync-dashboard\ProyInt\Plantilla_TRL4_Completado.docx",
        r"c:\asmasync-dashboard\ProyInt\Plantilla TRL4.docx"
    ]

    for doc_path in doc_files:
        if not os.path.exists(doc_path):
            continue
        try:
            doc = Document(doc_path)
            for p in list(doc.paragraphs):
                txt = p.text.strip()
                if "DIAGRAMA DE ARQUITECTURA CON LA INTERACCIÓN DE LOS USUARIOS" in txt or "Paso 4 - Visualización e Intervención Médica" in txt:
                    insert_image_after(p, target_img, "Figura 2. Diagrama de Interacción de los Usuarios de la Aplicación (AsmaSync)")
                    break

            doc.save(doc_path)
            print(f"Inserted User Interaction Diagram into {doc_path}")
        except PermissionError:
            print(f"File {doc_path} is locked by Word. Skipping save.")
        except Exception as e:
            print(f"Error in {doc_path}: {e}")

if __name__ == "__main__":
    insert_user_diagram_to_docx()
