import os
import docx
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_justified_indented_paragraph(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.5)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    return p

def add_bullet_point(doc, title, text):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run_title = p.add_run(title + ": ")
    run_title.bold = True
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(12)
    run_text = p.add_run(text)
    run_text.font.name = 'Arial'
    run_text.font.size = Pt(12)
    return p

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    run.font.name = 'Arial'
    return p

def main():
    doc = Document()

    # Define base styles
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(12)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("Reporte de Frameworks Seleccionados y Beneficios\nProyecto: AsmaSync")
    run_title.bold = True
    run_title.font.size = Pt(14)
    run_title.font.name = 'Arial'

    p_space = doc.add_paragraph()

    # Section 1
    add_heading(doc, "1. Reporte de Frameworks Seleccionados")
    
    add_justified_indented_paragraph(doc, "Para el desarrollo integral del sistema de monitoreo y predicción de crisis asmáticas (AsmaSync), se ha definido una arquitectura orientada a servicios utilizando tecnologías modernas. Se han seleccionado los siguientes frameworks fundamentales:")
    
    add_bullet_point(doc, "Angular (Web)", "Seleccionado para el desarrollo de la aplicación web (Dashboard principal). Su robustez y madurez permiten crear interfaces estructuradas, responsivas y modulares, ideales para el personal de salud y administradores del sistema.")
    add_bullet_point(doc, "Flutter (Móvil)", "Seleccionado para el desarrollo de la aplicación móvil. Al permitir la creación de aplicaciones multiplataforma, garantiza que los pacientes dispongan de una app unificada, rápida y accesible desde sus dispositivos personales para recibir notificaciones y monitorear su estado.")
    add_bullet_point(doc, "FastAPI (Backend)", "Seleccionado para el desarrollo del backend. Este framework permite una integración nativa y de alto rendimiento con los algoritmos predictivos construidos en Python, proveyendo endpoints rápidos y eficientes para el consumo de la web y el móvil.")

    p_space2 = doc.add_paragraph()

    # Section 2
    add_heading(doc, "2. Beneficios Directos en el Proyecto")

    add_justified_indented_paragraph(doc, "La selección de estos frameworks específicos aporta las siguientes ventajas clave directamente orientadas a los requerimientos funcionales y no funcionales del sistema:")

    add_bullet_point(doc, "Mantenibilidad y Escalabilidad (Angular)", "Al utilizar una arquitectura basada en componentes y tipado fuerte (TypeScript), el sistema web será mucho más tolerante a fallos y más fácil de escalar conforme se agreguen nuevos módulos médicos o de visualización.")
    add_bullet_point(doc, "Reducción de Tiempos y Costos (Flutter)", "Al contar con una única base de código para compilar de forma nativa a Android e iOS, se reduce significativamente el esfuerzo de programación y pruebas. Esto permite lanzar actualizaciones de la aplicación móvil de manera simultánea para todos los pacientes.")
    add_bullet_point(doc, "Alta Reactividad y Rendimiento (FastAPI y Flutter)", "La sincronización de datos en tiempo real requiere alta eficiencia. FastAPI ofrece un excelente rendimiento en el procesamiento del backend, mientras que Flutter garantiza transiciones y visualizaciones fluidas a 60fps en el dispositivo del usuario, lo cual es crítico para mostrar alertas de crisis respiratorias en tiempo real sin demoras.")
    add_bullet_point(doc, "Ecosistema Unificado y Seguro", "La integración entre el frontend estructurado de Angular, la versatilidad multiplataforma de Flutter y un backend robusto basado en Python facilita la implementación de capas de seguridad consistentes, autenticación compartida y protección rigurosa de la información sensible de los pacientes.")

    doc.save(r"c:\asmasync-dashboard\ProyInt\Reporte_Frameworks.docx")
    print("Document successfully created!")

if __name__ == '__main__':
    main()
