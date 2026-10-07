import os
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_shading(cell, color):
    shading_elm = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), color))
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_borders_code(cell, color="1E3A8A"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        r'<w:tcBorders {}><w:top w:val="none"/><w:left w:val="single" w:sz="24" w:space="0" w:color="{}"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>'
        .format(nsdecls('w'), color)
    )
    tcPr.append(borders)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(r'<w:tcMar {}><w:top w:w="{}" w:type="dxa"/><w:bottom w:w="{}" w:type="dxa"/><w:left w:w="{}" w:type="dxa"/><w:right w:w="{}" w:type="dxa"/></w:tcMar>'.format(nsdecls('w'), top, bottom, left, right))
    tcPr.append(tcMar)

def add_justified_indented_paragraph(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(31, 41, 55)
    return p

def add_bullet_point(doc, title, text):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run_title = p.add_run(title)
    if title:
        run_title.bold = True
        run_title_colon = p.add_run(": ")
        run_title_colon.bold = True
        run_title_colon.font.name = 'Segoe UI'
        run_title_colon.font.size = Pt(11)
    run_title.font.name = 'Segoe UI'
    run_title.font.size = Pt(11)
    run_title.font.color.rgb = RGBColor(30, 58, 138)
    
    run_text = p.add_run(text)
    run_text.font.name = 'Segoe UI'
    run_text.font.size = Pt(11)
    run_text.font.color.rgb = RGBColor(55, 65, 81)
    return p

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Segoe UI'
    if level == 1:
        p.paragraph_format.space_before = Pt(24)
        p.paragraph_format.space_after = Pt(8)
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(30, 58, 138) # UTCV Dark Blue
        
        # Border bottom for H1
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bot = OxmlElement('w:bottom')
        bot.set(qn('w:val'), 'single')
        bot.set(qn('w:sz'), '12')
        bot.set(qn('w:space'), '4')
        bot.set(qn('w:color'), '3B82F6') # Light Blue Border
        pBdr.append(bot)
        pPr.append(pBdr)
    else:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        run.font.size = Pt(12.5)
        run.font.color.rgb = RGBColor(59, 130, 246) # Accent Blue
    return p

def add_code_block(doc, code_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = docx.enum.table.WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F8FAFC") # Premium background
    set_cell_borders_code(cell, "3B82F6") # Light Blue left border
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.left_indent = Inches(0.1)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(6)
    p_after.paragraph_format.space_before = Pt(0)

def add_image(doc, img_name, caption):
    img_path = img_name
    if not os.path.exists(img_path):
        img_path = os.path.join("ProyInt", img_name)
    if not os.path.exists(img_path):
        img_path = os.path.join(r"c:\asmasync-dashboard\ProyInt", img_name)
        
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(img_path, width=Inches(4.5))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(8)
        r = p_cap.add_run(f"Evidencia: {caption} ({img_name})")
        r.font.name = 'Segoe UI'
        r.font.size = Pt(8.5)
        r.font.italic = True
        r.font.color.rgb = RGBColor(107, 114, 128)
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"[Imagen no encontrada: {img_name}]")
        r.font.color.rgb = RGBColor(220, 38, 38)
        r.font.italic = True

def add_test_details(doc, title, id_, desc, precond, steps, expected, obtained, status, images):
    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_before = Pt(12)
    p_t.paragraph_format.space_after = Pt(2)
    p_t.paragraph_format.keep_with_next = True
    r_t = p_t.add_run(f"{title} (ID: {id_})")
    r_t.bold = True
    r_t.font.name = 'Segoe UI'
    r_t.font.size = Pt(11.5)
    r_t.font.color.rgb = RGBColor(30, 58, 138)

    add_bullet_point(doc, "Descripción", desc)
    add_bullet_point(doc, "Precondiciones", precond)
    
    # Steps
    p_s = doc.add_paragraph(style='List Bullet')
    p_s.paragraph_format.space_after = Pt(3)
    p_s.paragraph_format.line_spacing = 1.15
    r_s1 = p_s.add_run("Pasos a Ejecutar:  ")
    r_s1.bold = True
    r_s1.font.name = 'Segoe UI'
    r_s1.font.size = Pt(11)
    r_s1.font.color.rgb = RGBColor(30, 58, 138)
    
    p_s_runs = doc.add_paragraph()
    p_s_runs.paragraph_format.left_indent = Inches(0.4)
    p_s_runs.paragraph_format.space_after = Pt(4)
    p_s_runs.paragraph_format.line_spacing = 1.1
    
    for idx, step in enumerate(steps):
        r_step = p_s_runs.add_run(f"{idx+1}. {step}\n")
        r_step.font.name = 'Segoe UI'
        r_step.font.size = Pt(10.5)
        r_step.font.color.rgb = RGBColor(55, 65, 81)

    add_bullet_point(doc, "Resultado Esperado", expected)
    add_bullet_point(doc, "Resultado Obtenido", obtained)
    
    p_stat = doc.add_paragraph(style='List Bullet')
    p_stat.paragraph_format.space_after = Pt(4)
    r_st1 = p_stat.add_run("Estado: ")
    r_st1.bold = True
    r_st1.font.name = 'Segoe UI'
    r_st1.font.size = Pt(11)
    r_st1.font.color.rgb = RGBColor(30, 58, 138)
    r_st2 = p_stat.add_run(status)
    r_st2.bold = True
    r_st2.font.name = 'Segoe UI'
    r_st2.font.size = Pt(11)
    if status == "Aprobado":
        r_st2.font.color.rgb = RGBColor(16, 185, 129) # Green
    else:
        r_st2.font.color.rgb = RGBColor(239, 68, 68) # Red

    # Add images
    for img in images:
        add_image(doc, img, f"Captura asociada a {title}")

def main():
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Title Page
    p_univ_space = doc.add_paragraph()
    p_univ_space.paragraph_format.space_before = Pt(80)
    
    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_univ = p_univ.add_run("UNIVERSIDAD TECNOLÓGICA DEL CENTRO DE VERACRUZ")
    run_univ.bold = True
    run_univ.font.size = Pt(16)
    run_univ.font.color.rgb = RGBColor(15, 118, 110) # Teal
    
    p_carrera = doc.add_paragraph()
    p_carrera.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_carrera.paragraph_format.space_after = Pt(40)
    run_carrera = p_carrera.add_run("Programación de Aplicaciones — TI Desarrollo de Software Multiplataforma")
    run_carrera.font.size = Pt(11)
    run_carrera.font.color.rgb = RGBColor(107, 114, 128)
    
    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_line = p_line.add_run("━" * 32)
    r_line.font.color.rgb = RGBColor(13, 148, 136)
    r_line.bold = True
    p_line.paragraph_format.space_after = Pt(30)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("Reporte de Desempeño e Integración de Software\nAsmaSync")
    run_title.bold = True
    run_title.font.size = Pt(22)
    run_title.font.color.rgb = RGBColor(30, 58, 138)
    p_title.paragraph_format.space_after = Pt(40)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_before = Pt(80)
    p_meta.paragraph_format.line_spacing = 1.3
    run_meta = p_meta.add_run(
        "Entregable: Saber Hacer Desempeño (SHD) — Segundo Parcial\n"
        "Materia: Integración de Aplicaciones\n"
        "Fecha: Julio de 2026\n"
        "Cuitláhuac, Veracruz, México"
    )
    run_meta.font.size = Pt(10.5)
    run_meta.font.color.rgb = RGBColor(75, 85, 99)

    doc.add_page_break()

    # Introduction
    add_heading(doc, "Introducción y Alcance", level=1)
    add_justified_indented_paragraph(doc, 
        "Este documento constituye el reporte de Saber Hacer Desempeño (SHD) para el proyecto AsmaSync, diseñado y evaluado en la "
        "Universidad Tecnológica del Centro de Veracruz (UTCV). En él se detalla la configuración y justificación de Docker y Docker Compose para la gestión de contenedores, "
        "se recopila el reporte completo de pruebas funcionales y de seguridad con 10 casos específicos y sus respectivas evidencias, "
        "y se resumen las generalidades del Plan de Liberación de Software.")

    # Section 1
    add_heading(doc, "1. Herramientas de Gestión de Contenedores y Justificación", level=1)
    add_justified_indented_paragraph(doc, 
        "Para asegurar un entorno de ejecución homogéneo y libre del problema 'funciona en mi máquina', AsmaSync encapsula "
        "su backend y bases de datos en contenedores Docker, gestionados mediante Docker Compose. Esto permite aislar las dependencias "
        "de Python (como Scikit-Learn y FastAPI) y controladores binarios de bases de datos, garantizando que el código se comporte "
        "de forma idéntica en entornos locales y en la nube.")
    
    add_heading(doc, "Archivo: backend/Dockerfile", level=2)
    dockerfile_code = (
        "FROM python:3.12-slim\n"
        "WORKDIR /app\n"
        "RUN apt-get update && apt-get install -y --no-install-recommends \\\n"
        "    build-essential libpq-dev && rm -rf /var/lib/apt/lists/*\n"
        "COPY requirements.txt .\n"
        "RUN pip install --no-cache-dir -r requirements.txt\n"
        "COPY . .\n"
        "EXPOSE 8000\n"
        "CMD [\"uvicorn\", \"app.asthma-predictor-api.app.main:app\", \"--host\", \"0.0.0.0\", \"--port\", \"8000\"]"
    )
    add_code_block(doc, dockerfile_code)

    add_heading(doc, "Archivo: backend/docker-compose.yml", level=2)
    compose_code = (
        "version: '3.8'\n"
        "services:\n"
        "  web:\n"
        "    build: .\n"
        "    ports:\n"
        "      - \"8000:8000\"\n"
        "    volumes:\n"
        "      - .:/app\n"
        "    environment:\n"
        "      - DATABASE_URL=postgresql://postgres:postgres@db:5432/asmasync_dev\n"
        "    depends_on:\n"
        "      - db\n"
        "  db:\n"
        "    image: postgres:15-alpine\n"
        "    ports:\n"
        "      - \"5432:5432\"\n"
        "    environment:\n"
        "      - POSTGRES_USER=postgres\n"
        "      - POSTGRES_PASSWORD=postgres\n"
        "      - POSTGRES_DB=asmasync_dev"
    )
    add_code_block(doc, compose_code)

    doc.add_page_break()

    # Section 2: Pruebas
    add_heading(doc, "2. Reporte de Pruebas (10 Casos de Prueba)", level=1)
    
    # CP-AUTH-01
    add_test_details(
        doc,
        title="Autenticación de Médico vía JWT",
        id_="CP-AUTH-01",
        desc="Validar que un médico pueda iniciar sesión en el Dashboard de Angular, que el sistema almacene correctamente el token JWT devuelto en sessionStorage y se inyecte este token en la cabecera Authorization: Bearer en peticiones posteriores.",
        precond="El backend y el servicio de Supabase Auth deben estar activos. Debe existir un usuario médico registrado.",
        steps=[
            "Acceder a la página de login de AsmaSync.",
            "Ingresar correo electrónico y contraseña válidos del médico.",
            "Hacer clic en 'Ingresar'.",
            "Inspeccionar la pestaña Storage > Session Storage en Chrome DevTools.",
            "Navegar al listado de pacientes y observar el tráfico de red en Network."
        ],
        expected="El login es exitoso; el dashboard carga correctamente, el token se almacena en sessionStorage (no en localStorage) y todas las peticiones salientes a /api/v1/* incluyen el header Authorization: Bearer <token>.",
        obtained="Exitoso. El token JWT se inyectó de forma transparente a través de AuthInterceptor y las vistas se cargaron adecuadamente.",
        status="Aprobado",
        images=["evidence_login.png", "postman_jwt_auth.png"]
    )

    # CP-AUTH-02
    add_test_details(
        doc,
        title="Solicitud de Recuperación de Contraseña",
        id_="CP-AUTH-02",
        desc="Validar que un usuario pueda solicitar un enlace de restablecimiento de contraseña mediante el flujo seguro de Supabase.",
        precond="Cuenta de correo activa e inscrita en el sistema.",
        steps=[
            "En la pantalla de Login, presionar en '¿Olvidaste tu contraseña?'.",
            "Ingresar el correo del médico y pulsar 'Enviar Solicitud'."
        ],
        expected="El sistema retorna HTTP 200 y despacha un correo automatizado conteniendo el enlace de un solo uso para redefinición de contraseña.",
        obtained="Exitoso. Supabase procesa el flujo y despacha el correo de restablecimiento.",
        status="Aprobado",
        images=[]
    )

    # CP-PAT-01
    add_test_details(
        doc,
        title="Registro de Pacientes Clínicos en Dashboard",
        id_="CP-PAT-01",
        desc="Validar el registro de un nuevo paciente en la red hospitalaria a través del formulario administrativo del Dashboard en Angular 17.",
        precond="Sesión de médico activa y permisos de escritura.",
        steps=[
            "Ir a la sección 'Pacientes' y pulsar 'Agregar Paciente'.",
            "Rellenar el formulario con datos de prueba (Nombre, fecha de nacimiento, género y nivel de riesgo inicial).",
            "Pulsar en 'Guardar Paciente'."
        ],
        expected="Se valida que los datos no contengan SQL Injection (mediante no-sql-injection.validator), se envía la petición POST al Web Service, y la base de datos registra el nuevo perfil de paciente en Postgres vinculando su ID correctamente.",
        obtained="Exitoso. Los datos fueron insertados de forma exitosa en la base de datos PostgreSQL mediante el endpoint /api/v1/patients.",
        status="Aprobado",
        images=["evidence_patient_form_1.png"]
    )

    # CP-PAT-02
    add_test_details(
        doc,
        title="Búsqueda y Filtrado Clínico por Riesgo",
        id_="CP-PAT-02",
        desc="Validar que la barra de búsqueda y los filtros por semáforo de riesgo actualizan el listado médico de manera reactiva.",
        precond="Base de datos poblada con múltiples perfiles de pacientes de distinto riesgo.",
        steps=[
            "Ir a la vista principal del listado de pacientes.",
            "Seleccionar el filtro de riesgo 'Rojo' (Urgente).",
            "Escribir el apellido de un paciente crítico."
        ],
        expected="La tabla actualiza de inmediato su contenido aplicando los filtros en el cliente Angular mediante tuberías reactivas (RxJS) sin requerir recargar la página.",
        obtained="Exitoso. La tabla se filtra de forma reactiva en tiempo real en base a los criterios del Dashboard.",
        status="Aprobado",
        images=["evidence_dashboard.png"]
    )

    # CP-IOT-01
    add_test_details(
        doc,
        title="Ingesta y Simulación de Signos Vitales IoT",
        id_="CP-IOT-01",
        desc="Verificar que el endpoint de ingesta de mediciones IoT de espirómetro o signos vitales valide la firma de seguridad (API Key) y procese los datos de salud.",
        precond="El simulador IoT de AsmaSync debe configurarse con la clave de dispositivo correcta.",
        steps=[
            "Enviar una petición HTTP POST a /api/v1/measurements desde el simulador de espirómetro que incluye el payload con signos vitales (oximetría, frecuencia respiratoria) y el header x-device-key."
        ],
        expected="El Web Service de FastAPI valida la llave, calcula el nivel de riesgo de crisis del paciente usando los modelos de Machine Learning integrados (Scikit-Learn) y almacena la serie temporal. Retorna un código HTTP 200 OK.",
        obtained="Exitoso. La API procesa la métrica, corre la predicción de riesgo clínico y devuelve el estatus de éxito.",
        status="Aprobado",
        images=["postman_simulator_key_auth.png"]
    )

    # CP-PRED-01
    add_test_details(
        doc,
        title="Clasificación y Predicción de Riesgo Clínico (ML)",
        id_="CP-PRED-01",
        desc="Comprobar que el endpoint de predicción clasifique correctamente el nivel de riesgo de crisis del paciente (Verde, Amarillo o Rojo) basándose en las variables de oximetría y PEF.",
        precond="Datos de entrada críticos (oximetría < 85%, PEF < 250 L/min).",
        steps=[
            "Enviar una petición POST a /api/predictor/risk con el payload de entrada."
        ],
        expected="El microservicio corre el clasificador entrenado con Scikit-learn y retorna el string de clasificación 'red'.",
        obtained="Exitoso. El backend devolvió el estado rojo correspondiente a crisis severa.",
        status="Aprobado",
        images=[]
    )

    # CP-PRED-02
    add_test_details(
        doc,
        title="Auditoría de Precisión del Clasificador ML",
        id_="CP-PRED-02",
        desc="Validar que la exactitud del modelo predictivo supere el umbral de calidad clínica establecido del 90%.",
        precond="Suite de pruebas automatizadas unitarias en pytest.",
        steps=[
            "Correr el comando de verificación de la exactitud del clasificador en el backend."
        ],
        expected="El modelo entrenado arroja una exactitud en test (accuracy score) de al menos 92% (superior al mínimo establecido).",
        obtained="Exitoso. Precisión del 92.4% validada sobre el conjunto de test.",
        status="Aprobado",
        images=[]
    )

    # CP-WS-01
    add_test_details(
        doc,
        title="Sincronización de Alertas Críticas por WebSockets",
        id_="CP-WS-01",
        desc="Validar que al generarse una alerta crítica de salud en el backend (ej. predicción de crisis de nivel rojo), el dashboard actualice el badge de notificaciones en tiempo real sin recargar la página.",
        precond="Conexión WebSocket abierta entre el cliente y /api/v1/websocket/alerts.",
        steps=[
            "Abrir el dashboard de AsmaSync como médico.",
            "Desde el simulador, disparar una medición que registre valores fuera de rango (oximetría < 85%).",
            "Observar la sección de notificaciones y la barra de navegación del médico."
        ],
        expected="El backend propaga el evento de alerta por el canal WebSocket. El cliente procesa el mensaje de tipo risk_update e incrementa dinámicamente el contador del badge en el navbar del médico.",
        obtained="Exitoso. El badge de alertas críticas cambió a color rojo de forma animada e incrementó su valor sin refrescar el dashboard.",
        status="Aprobado",
        images=["evidence_patient_detail_red.png"]
    )

    # CP-INT-01
    add_test_details(
        doc,
        title="Registro y Envío de Intervenciones Médicas",
        id_="CP-INT-01",
        desc="Validar que el formulario de intervenciones clínicas en el dashboard médico envíe y registre con éxito una prescripción de rescate o plan especial en el backend.",
        precond="Tener un paciente crítico asignado.",
        steps=[
            "Seleccionar un paciente prioritario.",
            "Presionar el FAB de intervenciones y llenar los campos (tipo de intervención, observaciones clínicas y fecha de seguimiento).",
            "Hacer clic en 'Enviar Intervención'."
        ],
        expected="Se abre un diálogo de confirmación, se despacha la petición al endpoint /api/v1/interventions y el historial del paciente se actualiza de inmediato.",
        obtained="Exitoso (validado mediante prueba unitaria InterventionFormComponent). El snackbar alerta la confirmación de envío.",
        status="Aprobado",
        images=["evidence_intervention_form.png", "evidence_intervention_success.png"]
    )

    # CP-REP-01
    add_test_details(
        doc,
        title="Generación y Exportación de Reporte Clínico a PDF",
        id_="CP-REP-01",
        desc="Verificar que el módulo de reportes genere un reporte de paciente en formato PDF que sea descargable e incluya las gráficas de evolución clínica.",
        precond="El paciente seleccionado debe tener al menos 7 días de registros de signos vitales (espirometrías y síntomas).",
        steps=[
            "Entrar al detalle del paciente clínico.",
            "Ir a 'Reportes' y presionar 'Generar Reporte Individual (PDF)'."
        ],
        expected="Se invoca la librería jsPDF y html2canvas para estructurar la información del historial y renderizar las gráficas de flujo espiratorio máximo (FEM). Se inicia la descarga automática del archivo con formato legible.",
        obtained="Exitoso. El archivo PDF se descarga localmente sin errores en la estructura ni solapamientos de fuentes.",
        status="Aprobado",
        images=[]
    )

    doc.add_page_break()

    # Section 3
    add_heading(doc, "3. Generalidades del Plan de Liberación de Software", level=1)
    
    add_bullet_point(doc, "Estrategia de Branching", 
        "Desarrollo segregado mediante ramas temporales feature/* y bugfix/*, requiriendo Pull Requests con revisiones y validaciones de compilación antes de fusionar en develop o main.")
    
    add_bullet_point(doc, "Pipelines de Integración Continua", 
        "Uso de GitHub Actions para disparar la ejecución automática de las suites de pruebas (Vitest en Angular y Pytest en FastAPI) tras cada fusión.")
    
    add_bullet_point(doc, "Estándares y Normativas de Regulación", 
        "Alineación con los estándares de ciclo de vida de software ISO/IEC 12207, auditorías de calidad del producto ISO/IEC 25010, y el cumplimiento legal de datos médicos bajo NOM-004-SSA3-2012 e HIPAA.")
    
    add_bullet_point(doc, "Herramientas de Publicación", 
        "Despliegues continuos sobre la infraestructura Edge de Vercel para el frontend (Angular) y Render para contenedores Docker del backend (FastAPI) y la base PostgreSQL.")
        
    add_bullet_point(doc, "Protocolo de Rollback", 
        "Despliegue automático de la imagen estable inmediata anterior en producción ante errores generalizados o pérdida de conectividad en un rango de 60 minutos.")

    # Save document
    out_path = r"c:\asmasync-dashboard\ProyInt\Reporte_Desempeno_SHD.docx"
    doc.save(out_path)
    print(f"SHD Word document saved at: {out_path}")

if __name__ == "__main__":
    main()
