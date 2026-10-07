import os
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def format_run(run, text, font_name="Montserrat", font_size=14, color_rgb=(51, 65, 85), bold=False, italic=False):
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(*color_rgb)
    run.bold = bold
    run.italic = italic

def format_p(p, text, font_size=14, color_rgb=(51, 65, 85), bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6):
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    p.text = ""
    run = p.add_run(text)
    format_run(run, text, font_size=font_size, color_rgb=color_rgb, bold=bold, italic=italic)
    return p

def update_layered_architecture():
    uploaded_img = r"C:\Users\penar\.gemini\antigravity\brain\ac6baca2-e01f-4959-b6b9-d862274c5e0a\.user_uploaded\media_1785955724526.png"
    target_img1 = r"c:\asmasync-dashboard\ProyInt\security_architecture.png"
    target_img2 = r"c:\asmasync-dashboard\ProyInt\layered_architecture.png"

    if os.path.exists(uploaded_img):
        shutil.copy(uploaded_img, target_img1)
        shutil.copy(uploaded_img, target_img2)
        print(f"Copied uploaded diagram to {target_img1} and {target_img2}")

    src_doc = r"c:\asmasync-dashboard\ProyInt\Plantilla_TRL4_Extenso.docx"
    if not os.path.exists(src_doc):
        src_doc = r"c:\asmasync-dashboard\ProyInt\Plantilla_TRL4_Completado.docx"

    doc = Document(src_doc)

    new_arch_intro = (
        "La arquitectura operacional de AsmaSync se fundamenta en un patrón de Arquitectura por Capas (N-Tier Layered Architecture) "
        "desacoplada, modular y altamente escalable, estructurada en cuatro niveles principales que garantizan la separación de responsabilidades, "
        "seguridad en el transporte de datos e inferencia de Machine Learning en tiempo real:"
    )

    p1 = (
        "• Capa de Presentación: Constituye la interfaz directa de interacción con los usuarios y dispositivos. Comprende los Dispositivos IoT "
        "(sensores biométricos, pulsiómetros y wearables para la transmisión de signos vitales) y las aplicaciones de usuario final desarrolladas en Flutter (Dart), "
        "las cuales abarcan tanto la Aplicación Móvil para pacientes y guardianes como el Dashboard Clínico Web para el personal de salud."
    )
    p2 = (
        "• Capa de Aplicación: Actúa como la puerta de entrada segura y controlada hacia el sistema. Incorpora un API Gateway encauzador, "
        "mecanismos de autenticación mediante firmas digitales JWT, políticas de Autorización estrictas basadas en roles (RBAC) y módulos de Rate Limiting "
        "para la protección contra ataques de denegación de servicio."
    )
    p3 = (
        "• Capa de Negocio: Concentra la lógica de procesamiento distribuido en contenedores Docker independientes. Incluye el módulo de Ingreso de Datos "
        "(desarrollado en Python) para la validación de biomarcadores; el bus de Mensajería distribuido con Apache Kafka para el enrutamiento asíncrono de eventos; "
        "el módulo de ML Predictivo (Python + Scikit-Learn) para la generación y cálculo de puntajes de riesgo; y el servicio de Notificaciones y Alertas (Python + FastAPI), "
        "el cual ejecuta la entrega de planes de acción e integra directamente con el servicio externo Firebase Cloud Messaging (FCM) para el despacho push a los guardianes."
    )
    p4 = (
        "• Capa de Persistencia: Proporciona almacenamiento híbrido optimizado. Emplea PostgreSQL como motor relacional primario para la gestión estructurada de expedientes, "
        "perfiles y tratamientos con políticas RLS, complementado con MongoDB para la persistencia no relacional de logs de eventos y series de tiempo biométricas."
    )

    for p in doc.paragraphs:
        txt = p.text.strip()
        if "La arquitectura operacional de AsmaSync se fundamenta" in txt or "comunicación cliente-servidor distribuida" in txt:
            p.text = new_arch_intro
            p_curr = p
            p_curr = format_p(doc.add_paragraph(), p1)
            p_curr = format_p(doc.add_paragraph(), p2)
            p_curr = format_p(doc.add_paragraph(), p3)
            p_curr = format_p(doc.add_paragraph(), p4)

        elif "Figura 1. Diagrama de Arquitectura" in txt:
            format_p(p, "Figura 1. Diagrama de Arquitectura por Capas (AsmaSync)", font_size=12, color_rgb=(100, 116, 139), italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    out_new = r"c:\asmasync-dashboard\ProyInt\Plantilla_TRL4_Capas.docx"
    doc.save(out_new)
    print(f"Saved Layered Architecture document to: {out_new}")

    # Intentar guardar en los nombres existentes si no están abiertos
    for path in [r"c:\asmasync-dashboard\ProyInt\Plantilla_TRL4_Extenso.docx", r"c:\asmasync-dashboard\ProyInt\Plantilla TRL4.docx"]:
        try:
            doc.save(path)
            print(f"Saved directly to {path}")
        except PermissionError:
            print(f"File {path} is currently locked by Word. Skipping save.")

if __name__ == "__main__":
    update_layered_architecture()
