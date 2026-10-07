import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def format_run(run, text, font_name="Montserrat", font_size=14, color_rgb=(51, 65, 85), bold=False, italic=False):
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(*color_rgb)
    run.bold = bold
    run.italic = italic

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def insert_image_after(p_ref, img_path, caption_text, width_inches=5.8):
    if not os.path.exists(img_path):
        print(f"Warning: Image {img_path} not found.")
        return p_ref

    p_img_elem = parse_xml(f'<w:p {nsdecls("w")}/>')
    p_ref._p.addnext(p_img_elem)
    doc_p_img = docx.text.paragraph.Paragraph(p_img_elem, p_ref._parent)
    doc_p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc_p_img.paragraph_format.space_before = Pt(14)
    doc_p_img.paragraph_format.space_after = Pt(4)
    
    run_img = doc_p_img.add_run()
    run_img.add_picture(img_path, width=Inches(width_inches))

    p_cap_elem = parse_xml(f'<w:p {nsdecls("w")}/>')
    p_img_elem.addnext(p_cap_elem)
    doc_p_cap = docx.text.paragraph.Paragraph(p_cap_elem, p_ref._parent)
    doc_p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc_p_cap.paragraph_format.space_after = Pt(16)
    
    run_cap = doc_p_cap.add_run(caption_text)
    format_run(run_cap, caption_text, font_size=12, color_rgb=(100, 116, 139), italic=True)
    return doc_p_cap

def fix_and_embed_trl4():
    doc_path = r"c:\asmasync-dashboard\ProyInt\Plantilla TRL4.docx"
    doc = Document(doc_path)
    proy_dir = r"c:\asmasync-dashboard\ProyInt"

    print("Cleaning leftover prompt instructions...")
    leftover_keywords = [
        "se deben considerar:",
        "se deberán elaborar las tablas del diccionario",
        "se debe incluir el diagrama relacional",
        "en esta fase se deberá mostrar las etapas",
        "por ejemplo, se puede utilizar el siguiente proceso",
        "se deberá evidenciar:",
        "se deberá mostrar la evidencia del manual del sistema",
        "directorios y archivos (se deben describir",
        "se debe describir cada objeto",
        "mostrar los objetos que interactúan",
        "se deben describir ejemplos de las url",
        "se deberá mostrar evidencia del manual de usuario aplicando"
    ]

    for p in doc.paragraphs:
        t_low = p.text.strip().lower()
        if any(kw in t_low for kw in leftover_keywords):
            p.text = ""

    print("Embedding all 7 physical images in proper document locations...")
    
    # 1. Diagrama de Arquitectura en 1.4
    for p in doc.paragraphs:
        txt = p.text.strip()
        if "1.4 Funcionalidad" in txt or "La arquitectura operacional de AsmaSync" in txt:
            insert_image_after(p, os.path.join(proy_dir, "security_architecture.png"), "Figura 1. Diagrama de Arquitectura de Seguridad y Comunicación Cliente-Servidor (AsmaSync)")
            break

    # 2-6. Evidencias del Manual del Usuario en 2.2
    for p in doc.paragraphs:
        txt = p.text.strip()
        if "Paso 4 - Despacho de Alertas" in txt:
            insert_image_after(p, os.path.join(proy_dir, "evidence_dashboard.png"), "Figura 2. Dashboard Clínico Web con Semaforización de Riesgo en Tiempo Real")
            insert_image_after(p, os.path.join(proy_dir, "evidence_patient_detail_red.png"), "Figura 3. Ficha de Expediente de Paciente en Riesgo Alto (Nivel Rojo)")
        elif "Paso 5 - Intervención Médica" in txt:
            insert_image_after(p, os.path.join(proy_dir, "evidence_intervention_form.png"), "Figura 4. Formulario de Registro de Intervención Médica Preventiva")
            insert_image_after(p, os.path.join(proy_dir, "evidence_intervention_success.png"), "Figura 5. Confirmación de Intervención Clínica Registrada en Supabase")
        elif "Paso 6 - Reporte PDF" in txt:
            insert_image_after(p, os.path.join(proy_dir, "evidence_pdf_report.png"), "Figura 6. Generación del Expediente Clínico Consolidado en Formato PDF")

    # 7. Postman REST API Inferencia en 2.6
    for p in doc.paragraphs:
        txt = p.text.strip()
        if "Resultados Evaluativos del Modelo Random Forest Classifier:" in txt:
            insert_image_after(p, os.path.join(proy_dir, "postman_api_key_auth.png"), "Figura 7. Validación de Inferencia de la API REST mediante Postman en Laboratorio")
            break

    out1 = r"c:\asmasync-dashboard\ProyInt\Plantilla TRL4.docx"
    out2 = r"c:\asmasync-dashboard\ProyInt\Plantilla_TRL4_Completado.docx"
    doc.save(out1)
    doc.save(out2)
    print(f"Final embed & cleanup complete! Saved to {out1} and {out2}")

if __name__ == "__main__":
    fix_and_embed_trl4()
