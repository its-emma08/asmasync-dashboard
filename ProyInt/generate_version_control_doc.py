import os
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_shading(cell, color):
    shading_elm = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), color))
    cell._tc.get_or_add_tcPr().append(shading_elm)

def add_placeholder(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.0)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F3F4F6") # Light gray
    
    # Set row height
    tr = table.rows[0]._tr
    trHeight = parse_xml(r'<w:trHeight xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="1440" w:hRule="exact"/>')
    tr.get_or_add_trPr().append(trHeight)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Add vertical spacing
    p.paragraph_format.space_before = Pt(30)
    run = p.add_run(f"📷 [ ESPACIO PARA CAPTURA DE PANTALLA ]\n{text}")
    run.font.name = 'Arial'
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(107, 114, 128)
    
    p_space = doc.add_paragraph() # Spacing after table
    p_space.paragraph_format.space_after = Pt(10)

def add_justified_indented_paragraph(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.5)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    return p

def add_bullet_point(doc, title, text):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run_title = p.add_run(title)
    if title:
        run_title.bold = True
        run_title_colon = p.add_run(": ")
        run_title_colon.bold = True
        run_title_colon.font.name = 'Arial'
        run_title_colon.font.size = Pt(12)
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(12)
    run_text = p.add_run(text)
    run_text.font.name = 'Arial'
    run_text.font.size = Pt(12)
    return p

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12 if level > 1 else 14)
    run.font.name = 'Arial'
    return p

def main():
    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(12)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("Control de Versiones y Flujo de Trabajo\nProyecto Integrador: AsmaSync")
    run_title.bold = True
    run_title.font.size = Pt(16)
    
    doc.add_paragraph()

    # 1. Justificación
    add_heading(doc, "1. Justificación de las Plataformas y Herramientas de Versionamiento")
    add_justified_indented_paragraph(doc, "Para el desarrollo del proyecto AsmaSync, se ha seleccionado Git como sistema de control de versiones distribuido y GitHub como plataforma de alojamiento y colaboración remota. La elección de Git se justifica por su capacidad nativa para gestionar el historial de cambios, facilitar el trabajo asíncrono y evitar la pérdida de código fuente. GitHub aporta valor significativo al ecosistema mediante sus herramientas colaborativas, el sistema de revisión de código a través de Pull Requests y sus robustas opciones de protección de ramas.")
    add_placeholder(doc, "Captura sugerida: Interfaz principal del repositorio en GitHub o de Git local.")

    # 2. Flujo de Trabajo
    add_heading(doc, "2. Flujo de Trabajo del Control de Versiones")
    add_justified_indented_paragraph(doc, "Se ha adoptado el flujo de trabajo 'Feature Branch Workflow' (basado en un Git Flow simplificado). En este modelo, todo el desarrollo de nuevas características, refactorizaciones o correcciones de errores se aísla de inmediato en ramas independientes. Una vez que la característica se completa, se realiza un Pull Request hacia la rama de desarrollo (develop) para su respectiva revisión por parte de otros miembros del equipo. Cuando el estado de desarrollo es estable y validado, los cambios se fusionan a la rama principal (main) para su eventual despliegue.")
    add_placeholder(doc, "Captura sugerida: Gráfico de red (Network Graph) de GitHub o historial de commits.")

    # 3. Descripción de Ramas
    add_heading(doc, "3. Descripción de Todas las Ramas a Utilizar")
    add_justified_indented_paragraph(doc, "El esquema de ramas estructurado para todo el equipo de desarrollo del proyecto es el siguiente:")
    add_bullet_point(doc, "main (principal)", "Contiene el código estable y funcional que se encuentra listo para el entorno de producción. Solo se actualiza mediante fusiones (merges) aprobadas y verificadas provenientes desde la rama de desarrollo.")
    add_bullet_point(doc, "develop (desarrollo)", "Rama base donde convergen todas las características desarrolladas por el equipo. Representa el código en etapa de pre-producción, 'staging' o integración continua.")
    add_bullet_point(doc, "feature/[nombre] o feature/[tarea]", "Ramas temporales utilizadas por los diferentes miembros del equipo para desarrollar nuevas funcionalidades sin afectar el código base común. Ejemplos de uso: 'feature/login-paciente', 'feature/emma-dashboard' o 'feature/auth'.")
    add_bullet_point(doc, "bugfix/[nombre-error]", "Ramas creadas con el propósito de aislar la corrección de errores puntuales encontrados en el entorno de desarrollo o pruebas.")
    add_placeholder(doc, "Captura sugerida: Lista de ramas activas en GitHub (Branch list).")

    # 4. Configuración
    add_heading(doc, "4. Configuración de las Plataformas y Herramientas de Versionamiento")
    add_justified_indented_paragraph(doc, "La configuración del entorno requirió establecer parámetros locales en las terminales de los desarrolladores mediante los comandos 'git config user.name' y 'git config user.email' para identificar de manera inequívoca la autoría de cada 'commit'. A nivel de la plataforma GitHub, se configuraron Reglas de Protección de Ramas (Branch Protection Rules) sobre 'main' y 'develop', las cuales exigen al menos la aprobación de un revisor en los Pull Requests antes de permitir una fusión, previniendo así la alteración accidental del código.")
    add_placeholder(doc, "Captura sugerida: Pantalla de 'Settings > Branches' en GitHub mostrando las reglas de protección.")

    # 5. Enlace del Repositorio
    add_heading(doc, "5. Enlace del Repositorio en Funcionamiento")
    add_justified_indented_paragraph(doc, "El repositorio se encuentra activo y público (o privado colaborativo), documentando todo el historial del proyecto. La trazabilidad completa del trabajo de cada integrante es fácilmente identificable mediante la lista de contribuciones individuales ('Commits' y 'Pull Requests').")
    
    p_link = doc.add_paragraph()
    p_link.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_link.paragraph_format.first_line_indent = Inches(0.5)
    run_url_title = p_link.add_run("URL del Repositorio: ")
    run_url_title.bold = True
    run_url_title.font.name = 'Arial'
    run_url_title.font.size = Pt(12)
    
    run_url = p_link.add_run("https://github.com/its-emma08/asmasync-dashboard")
    run_url.font.name = 'Arial'
    run_url.font.size = Pt(12)
    run_url.font.color.rgb = RGBColor(5, 99, 193)
    run_url.font.underline = True

    add_placeholder(doc, "Captura sugerida: Pestaña 'Insights > Contributors' en GitHub donde se vea a los miembros del equipo.")

    doc.save(r"c:\asmasync-dashboard\ProyInt\Control_Versiones_Proyecto.docx")
    print("Document successfully created!")

if __name__ == '__main__':
    main()
