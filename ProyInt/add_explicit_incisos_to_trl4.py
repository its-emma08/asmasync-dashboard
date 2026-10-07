import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def format_run(run, text, font_name="Montserrat", font_size=14, color_rgb=(51, 65, 85), bold=False, italic=False):
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(*color_rgb)
    run.bold = bold
    run.italic = italic

def format_p(p, text, font_size=14, color_rgb=(51, 65, 85), bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6):
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    p.text = ""
    run = p.add_run(text)
    format_run(run, text, font_size=font_size, color_rgb=color_rgb, bold=bold, italic=italic)
    return p

def insert_h3_after(p_ref, text):
    new_p_elem = parse_xml(f'<w:p {nsdecls("w")}/>')
    p_ref._p.addnext(new_p_elem)
    doc_p = docx.text.paragraph.Paragraph(new_p_elem, p_ref._parent)
    doc_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc_p.paragraph_format.space_before = Pt(14)
    doc_p.paragraph_format.space_after = Pt(6)
    doc_p.paragraph_format.keep_with_next = True
    run = doc_p.add_run(text)
    format_run(run, text, font_size=14, color_rgb=(30, 41, 59), bold=True)
    return doc_p

def insert_p_after(p_ref, text, font_size=14, color_rgb=(51, 65, 85), bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6):
    new_p_elem = parse_xml(f'<w:p {nsdecls("w")}/>')
    p_ref._p.addnext(new_p_elem)
    doc_p = docx.text.paragraph.Paragraph(new_p_elem, p_ref._parent)
    doc_p.alignment = align
    doc_p.paragraph_format.space_after = Pt(space_after)
    doc_p.paragraph_format.line_spacing = 1.25
    run = doc_p.add_run(text)
    format_run(run, text, font_size=font_size, color_rgb=color_rgb, bold=bold, italic=italic)
    return doc_p

def insert_bullet_after(p_ref, title, desc=""):
    new_p_elem = parse_xml(f'<w:p {nsdecls("w")}/>')
    p_ref._p.addnext(new_p_elem)
    doc_p = docx.text.paragraph.Paragraph(new_p_elem, p_ref._parent)
    doc_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc_p.paragraph_format.space_after = Pt(6)
    doc_p.paragraph_format.left_indent = Inches(0.25)
    doc_p.paragraph_format.line_spacing = 1.25

    rb = doc_p.add_run("• ")
    rb.font.name = "Montserrat"
    rb.font.size = Pt(14)
    rb.font.color.rgb = RGBColor(2, 132, 199)
    rb.bold = True

    rt = doc_p.add_run(title)
    rt.font.name = "Montserrat"
    rt.font.size = Pt(14)
    rt.font.color.rgb = RGBColor(15, 23, 42)
    rt.bold = True

    if desc:
        rd = doc_p.add_run(f": {desc}")
        rd.font.name = "Montserrat"
        rd.font.size = Pt(14)
        rd.font.color.rgb = RGBColor(51, 65, 85)
    return doc_p

def update_trl4_incisos():
    doc_path = r"c:\asmasync-dashboard\ProyInt\Plantilla TRL4.docx"
    doc = Document(doc_path)

    # Iterar párrafos para estructurar incisos explícitos a), b), c), d), e), f), g), h)
    for p in list(doc.paragraphs):
        txt = p.text.strip()

        # 2.4 Exploración
        if "2.4 Exploración" in txt or "2.4 Exploraci" in txt:
            curr = p
            curr = insert_h3_after(curr, "a) Establecimiento de Actores")
            curr = insert_p_after(curr, "Se establecen los actores principales que participan en el desarrollo e interacción tecnológica:")
            curr = insert_bullet_after(curr, "Paciente", "Usuario final asmático que registra biomarcadores cotidianos y sigue indicaciones.")
            curr = insert_bullet_after(curr, "Guardián (Familiar/Tutor)", "Usuario responsable que recibe alertas de emergencia push FCM.")
            curr = insert_bullet_after(curr, "Doctor / Personal de Salud", "Profesional médico que monitorea el Dashboard, analiza tendencias y emite intervenciones.")
            curr = insert_bullet_after(curr, "Administrador del Sistema", "Personal de TI a cargo del control de cuentas, roles RBAC y servidores.")

            curr = insert_h3_after(curr, "b) Definición del Alcance")
            curr = insert_p_after(curr, "El sistema abarca el monitoreo preventivo continuo, la ingesta en tiempo real de signos vitales, la inferencia de riesgo con IA (Random Forest), el despacho automático de notificaciones de alerta y la exportación de expedientes en PDF.")

            curr = insert_h3_after(curr, "c) Requerimientos Funcionales y No Funcionales")
            curr = insert_p_after(curr, "Se especifican los requerimientos funcionales (RF01 a RF08) y no funcionales (RNF01 a RNF02) para el correcto desempeño de la plataforma.")

            curr = insert_h3_after(curr, "d) Diagramas de Casos de Uso")
            curr = insert_p_after(curr, "Los casos de uso definen la interacción de cada actor con la plataforma (Ingreso de síntomas, Consulta de semáforo, Registro de intervenciones y Exportación de reportes).")

            curr = insert_h3_after(curr, "e) Tabla de Datos con las Principales Variables a Trabajar")
            curr = insert_p_after(curr, "Se presenta el catálogo de variables clínicas y ambientales analizadas por el sistema (frecuencia cardíaca, FR, SpO2, uso de inhalador, sibilancias, temperatura y AQI).")

            curr = insert_h3_after(curr, "f) Procesos (Diagrama de Actividades)")
            curr = insert_p_after(curr, "El diagrama de actividades describe el flujo secuencial desde la captura en el cliente móvil, validación de JWT, vectorización de variables, inferencia en predict.py, inserción en Supabase y refresco del Dashboard.")

        # 2.5 Iniciación
        elif "2.5 Iniciación" in txt or "2.5 Iniciaci" in txt:
            curr = p
            curr = insert_h3_after(curr, "a) Establecimiento de Recursos Físicos")
            curr = insert_p_after(curr, "Se definen los recursos físicos y de infraestructura necesarios:")
            curr = insert_bullet_after(curr, "Dispositivos Móviles", "Smartphones Android 8.0+ / iOS 12.0+.")
            curr = insert_bullet_after(curr, "Servidor Backend", "Instancia Web Service en Render (Python 3.11 / Uvicorn).")
            curr = insert_bullet_after(curr, "Base de Datos y Auth", "Supabase BaaS (PostgreSQL 15 con extensiones RLS).")
            curr = insert_bullet_after(curr, "Servicio de Notificaciones", "Infraestructura Google Firebase Cloud Messaging (FCM).")
            curr = insert_bullet_after(curr, "Estación Médica", "Laptops/PC con navegador web moderno para el Dashboard.")

            curr = insert_h3_after(curr, "b) Establecimiento de Comunicación")
            curr = insert_p_after(curr, "La comunicación cliente-servidor utiliza servicios RESTful sobre HTTPS (TLS 1.3) con tokens JWT Bearer y payloads JSON.")

            curr = insert_h3_after(curr, "c) Modelado de Datos (Diccionario de Datos y ERD)")
            curr = insert_p_after(curr, "Se incluyen las tablas del diccionario de datos (profiles, health_records, interventions) con sus tipos, tamaños, nulos y descripciones, acompañadas del diagrama relacional ERD.")

            curr = insert_h3_after(curr, "d) Modelado de Componentes")
            curr = insert_p_after(curr, "El desarrollo se conforma de 4 módulos principales: Módulo de Autenticación, Motor de Inferencia ML, Servicio de Notificaciones FCM y Generador PDF.")

        # 2.6 Producción
        elif "2.6 Producción" in txt or "2.6 Producci" in txt:
            curr = p
            curr = insert_p_after(curr, "Proceso Cíclico de Metodología Ágil (Scrum):")
            curr = insert_bullet_after(curr, "a) Planeación de la tarea", "Definición del backlog, diseño del esquema de base de datos e ingesta de variables.")
            curr = insert_bullet_after(curr, "b) Inicio de la tarea", "Configuración del entorno virtual Python, repositorios Git y proyecto Angular/Flutter.")
            curr = insert_bullet_after(curr, "c) Desarrollo de la tarea", "Entrenamiento del modelo Random Forest, construcción de endpoints FastAPI y vistas de UI.")
            curr = insert_bullet_after(curr, "d) Pruebas", "Pruebas de comunicación REST, validación de inferencia en milisegundos y recepción push FCM.")

            curr = insert_h3_after(curr, "a) Diseño de Interfaces (Mockups)")
            curr = insert_p_after(curr, "Se presentan los diseños de interfaz del Dashboard Web y de la App Móvil con su tabla de propiedades y capturas visuales.")

            curr = insert_h3_after(curr, "b) Análisis de Resultados de Diferentes Algoritmos (Machine Learning)")
            curr = insert_p_after(curr, "Se evaluaron 4 clasificadores (Regresión Logística, SVM, XGBoost y Random Forest), seleccionando Random Forest por alcanzar un 89.4% de exactitud y 87.2% de sensibilidad.")

            curr = insert_h3_after(curr, "c) Pruebas de los Estudios (Clasificación/Predicción)")
            curr = insert_p_after(curr, "Se presentan los resultados de las pruebas de laboratorio sobre el conjunto de datos de validación.")

            curr = insert_h3_after(curr, "d) Desarrollo de API REST")
            curr = insert_p_after(curr, "Especificación formal de los controladores y endpoints de la API (auth, predict, patients, interventions, reports).")

            curr = insert_h3_after(curr, "e) Configuración de la Base de Datos")
            curr = insert_p_after(curr, "Scripts de creación de tablas, índices de rendimiento y políticas RLS en Supabase.")

            curr = insert_h3_after(curr, "f) Repositorio de Control de Versiones")
            curr = insert_p_after(curr, "Estructura de ramas Git en GitHub y pipelines de despliegue automático (auto-deploy) en Render.")

            curr = insert_h3_after(curr, "g) Código de las Principales Funciones Documentado")
            curr = insert_p_after(curr, "Fragmentos documentados de las funciones clave (ej. función de predicción predict_risk() en predict.py).")

            curr = insert_h3_after(curr, "h) Desarrollo y Pruebas de API con Postman")
            curr = insert_p_after(curr, "Evidencia de peticiones exitosas a la API mediante Postman en entorno de laboratorio.")

        # 2.7 Estabilización
        elif "2.7 Estabilización" in txt or "2.7 Estabilizaci" in txt:
            curr = p
            curr = insert_p_after(curr, "Manual del Sistema Simplificado:")
            curr = insert_bullet_after(curr, "a) Entidades del proyecto", "Tablas profiles, health_records, interventions y guardians_patients.")
            curr = insert_bullet_after(curr, "b) Métodos utilizables de la API", "Operaciones HTTP POST, GET y sus permisos asociados.")
            curr = insert_bullet_after(curr, "c) Directorios y archivos", "Estructura de carpetas backend/, frontend/ y asthmaapp/.")
            curr = insert_bullet_after(curr, "d) Clases, métodos y objetos del sistema", "Modelos Pydantic (PredictRequest, PredictResponse) y SQLAlchemy.")
            curr = insert_bullet_after(curr, "e) Dispositivos", "Objetos de interfaz con pulsiómetros Bluetooth, sensores móviles y APIs de calidad del aire.")
            curr = insert_bullet_after(curr, "f) URL y datos en formato JSON", "Ejemplos completos de peticiones y respuestas HTTP 200 OK.")

    out1 = r"c:\asmasync-dashboard\ProyInt\Plantilla TRL4.docx"
    out2 = r"c:\asmasync-dashboard\ProyInt\Plantilla_TRL4_Completado.docx"
    doc.save(out1)
    doc.save(out2)
    print("Added explicit lettered incisos a)-h) to both template files successfully!")

if __name__ == "__main__":
    update_trl4_incisos()
